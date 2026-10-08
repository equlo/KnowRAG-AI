"""
loaders.py - Step 1 of RAG ingestion: turn an uploaded FILE into plain TEXT.

Each supported format has its own small function. They all return a list of
`Page` objects so that PDFs can remember which page each piece of text came
from (useful for citations like "handbook.pdf, page 4").
"""

# `io.BytesIO` wraps raw bytes so libraries that expect a file object can read them.
import io

# `dataclass` auto-generates __init__/__repr__ for simple data-holding classes.
from dataclasses import dataclass

# `PurePath` lets us read a file name's extension without touching the disk.
from pathlib import PurePath

# PdfReader parses PDF files.
from pypdf import PdfReader

# `Document` (renamed to DocxDocument to avoid confusion) parses .docx files.
from docx import Document as DocxDocument


# The set of file extensions this application knows how to read.
SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx"}


# A custom error type so the API layer can turn it into a friendly 400 response.
class UnsupportedFileError(ValueError):
    """Raised when a file cannot be turned into text."""


# A tiny container: one page (or one whole file) of extracted text.
# frozen=True makes instances immutable, which avoids accidental edits.
@dataclass(frozen=True)
class Page:
    # The extracted text.
    text: str
    # 1-based page number for PDFs; None when the format has no pages.
    page: int | None = None


def load_document(filename: str, data: bytes) -> list[Page]:
    """Dispatch to the right reader based on the file extension."""
    # Get the extension in lower case, e.g. "Report.PDF" -> ".pdf".
    extension = PurePath(filename).suffix.lower()
    # PDF files are read page by page.
    if extension == ".pdf":
        pages = _load_pdf(data)
    # Word files are read paragraph by paragraph.
    elif extension == ".docx":
        pages = _load_docx(data)
    # Plain text and Markdown are simply decoded from bytes to a string.
    elif extension in {".txt", ".md"}:
        pages = [Page(text=_decode_text(data))]
    # Anything else is rejected with a clear message.
    else:
        allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedFileError(f"Unsupported file type '{extension}'. Allowed: {allowed}")

    # Drop pages that contain only whitespace (e.g. blank PDF pages).
    pages = [p for p in pages if p.text.strip()]
    # If nothing is left, the file is probably a scanned image or empty.
    if not pages:
        raise UnsupportedFileError(
            f"No text could be extracted from '{filename}'. "
            "Scanned PDFs need OCR before they can be indexed."
        )
    # Hand the list of pages to the next stage (chunking).
    return pages


def _decode_text(data: bytes) -> str:
    """Convert raw bytes to text, trying the most common encodings."""
    # Try UTF-8 first (the modern standard); "utf-8-sig" also strips a BOM marker.
    try:
        return data.decode("utf-8-sig")
    # If that fails, fall back to Latin-1, which can decode ANY byte sequence.
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _load_pdf(data: bytes) -> list[Page]:
    """Extract text from every page of a PDF."""
    # Open the PDF from memory (no temporary file needed).
    try:
        reader = PdfReader(io.BytesIO(data))
    # pypdf raises various errors for corrupt files; report them uniformly.
    except Exception as exc:
        raise UnsupportedFileError(f"Could not read PDF: {exc}") from exc
    # Build one Page per PDF page; enumerate(..., start=1) gives human page numbers.
    # `extract_text()` can return None for image-only pages, so we use `or ""`.
    return [Page(text=page.extract_text() or "", page=i) for i, page in enumerate(reader.pages, start=1)]


def _load_docx(data: bytes) -> list[Page]:
    """Extract text from a Word document."""
    # Open the .docx from memory.
    try:
        document = DocxDocument(io.BytesIO(data))
    # A renamed or corrupt file will fail to open.
    except Exception as exc:
        raise UnsupportedFileError(f"Could not read DOCX: {exc}") from exc
    # Collect the text of every paragraph.
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    # Tables are stored separately from paragraphs, so read each cell too.
    for table in document.tables:
        # Each row becomes one line, with cells separated by " | ".
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    # Word has no fixed pages, so the whole document is one "page" with page=None.
    # Blank lines between paragraphs let the chunker detect paragraph boundaries.
    return [Page(text="\n\n".join(paragraphs))]
